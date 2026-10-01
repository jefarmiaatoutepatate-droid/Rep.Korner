"""Convertit Yellow_Jack_Comptabilite.numbers (Apple Numbers) en .xlsx lisible sous Windows / Excel / LibreOffice / Google Sheets.

Usage : pip install numbers-parser openpyxl
        python numbers_to_xlsx.py chemin/vers/fichier.numbers [sortie.xlsx]

Les valeurs et la mise en forme sont recopiees depuis le fichier Numbers. Les formules qui
pointent vers un autre onglet (que numbers-parser ne sait pas traduire) sont reecrites a la main
en syntaxe Excel.
"""
import re
import sys

from numbers_parser import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MONEY = '#,##0 "$";-#,##0 "$";"-"'
PCT = "0.0%"
DATE = "dd/mm/yyyy"
LAST = {"Ventes": 201, "Depenses": 201, "Salaires": 200, "Stock": 203}  # derniere ligne de saisie


def rgb(c):
    return f"{c.r:02X}{c.g:02X}{c.b:02X}" if c else None


def excel_formula(sheet, row, col, f):
    """row/col 1-based. Renvoie la formule Excel (sans '=')."""
    r = row
    if sheet == "Tableau de bord":
        return {
            (6, 2): "TODAY()",
            (9, 2): f"SUM(Ventes!G5:G{LAST['Ventes']})",
            (12, 2): f"SUM(Depenses!D5:D{LAST['Depenses']})",
            (13, 2): f"SUM(Salaires!F5:F{LAST['Salaires']})",
        }.get((row, col)) or convert(f)
    if sheet == "Resume mensuel" and 7 <= r <= 18 and col in (2, 3, 4):
        m = r - 6
        src, val = {2: ("Ventes", "G"), 3: ("Depenses", "D"), 4: ("Salaires", "F")}[col]
        n = LAST[src]
        return (f'SUMIFS({src}!{val}5:{val}{n},{src}!A5:A{n},">="&DATE($B$4,{m},1),'
                f'{src}!A5:A{n},"<"&DATE($B$4,{m}+1,1))')
    if "::" in f:
        raise ValueError(f"Reference inter-onglet non geree : {sheet} {row},{col} {f}")
    return convert(f)


def convert(f):
    return f.replace("×", "*").replace("÷", "/")


# Colonnes calculees (police noire dans Numbers) : on y remet toujours la formule,
# protegee par IFERROR pour ne pas afficher d'erreur tant que des "X" restent en saisie.
COMPUTED = {
    "Ventes": (7, lambda r: f'IFERROR(IF(E{r}="","",E{r}*F{r}),"")'),
    "Salaires": (6, lambda r: f'IFERROR(IF(D{r}="","",D{r}*E{r}),"")'),
    "Stock": (7, lambda r: f'IFERROR(IF(A{r}="","",F{r}-E{r}),"")'),
}
NUMFMT = {
    "Ventes": {1: DATE, 6: MONEY, 7: MONEY},
    "Depenses": {1: DATE, 4: MONEY},
    "Salaires": {1: DATE, 5: MONEY, 6: MONEY},
    "Stock": {5: MONEY, 6: MONEY, 7: MONEY},
    "Resume mensuel": {2: MONEY, 3: MONEY, 4: MONEY, 5: MONEY, 6: MONEY, 7: PCT},
}


def main(src, dst):
    doc = Document(src)
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in doc.sheets:
        t = sheet.tables[0]
        ws = wb.create_sheet(sheet.name)
        for c in range(t.num_cols):
            ws.column_dimensions[get_column_letter(c + 1)].width = max(10, t.col_width(c) / 6.5)
        for m in t.merge_ranges:
            ws.merge_cells(m)
        for r in range(t.num_rows):
            for c in range(t.num_cols):
                cell = t.cell(r, c)
                row, col = r + 1, c + 1
                if cell.is_formula:
                    value = "=" + excel_formula(sheet.name, row, col, cell.formula)
                elif cell.value in (None, ""):
                    value = None
                else:
                    value = cell.value
                if sheet.name in COMPUTED and row >= 5 and row <= LAST[sheet.name]:
                    ccol, fn = COMPUTED[sheet.name]
                    if col == ccol:
                        value = "=" + fn(row)
                x = ws.cell(row, col)
                if value is not None:
                    x.value = value
                st = cell.style
                if st:
                    x.font = Font(name=st.font_name, size=st.font_size, bold=st.bold,
                                  italic=st.italic, color=rgb(st.font_color))
                    if st.bg_color:
                        x.fill = PatternFill("solid", fgColor=rgb(st.bg_color))
                fmt = NUMFMT.get(sheet.name, {}).get(col)
                if fmt and row >= 5:
                    x.number_format = fmt
        ws.freeze_panes = "A5" if sheet.name in LAST else None

    # Tableau de bord : formats + indicateurs non renseignes dans le fichier d'origine
    ws = wb["Tableau de bord"]
    ws["B6"].number_format = "mmmm yyyy"
    for a in ("B9", "B12", "B13", "B14", "B17", "B21"):
        ws[a].number_format = MONEY
    ws["B18"].number_format = PCT
    n = LAST["Stock"]
    ws["B21"] = f"=SUMPRODUCT(Stock!C5:C{n},Stock!E5:E{n})"
    ws["B22"] = (f"=SUMPRODUCT(ISNUMBER(Stock!C5:C{n})*ISNUMBER(Stock!D5:D{n})"
                 f"*(Stock!C5:C{n}<=Stock!D5:D{n}))")
    for a in ("B9", "B12", "B13", "B14", "B17", "B18", "B21", "B22"):
        ws[a].alignment = Alignment(horizontal="right")
    # Salaires : ligne TOTAL vide dans l'original
    s = wb["Salaires"]
    s["F201"] = f"=SUM(F5:F{LAST['Salaires']})"
    s["F201"].number_format = MONEY
    s["E201"].font = s["F201"].font = Font(name="Arial", size=11, bold=True)
    wb.save(dst)
    print("OK ->", dst)


if __name__ == "__main__":
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else re.sub(r"\.numbers$", "", src) + ".xlsx"
    main(src, dst)
